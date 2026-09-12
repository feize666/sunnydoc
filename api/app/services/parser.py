"""文档解析服务：多格式提取文本 + zip 媒体提取 + 扫描件 OCR"""
from __future__ import annotations

import io
import os
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from app.core.config import TEXT_EXTS

# 媒体扩展名集合
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".ico"}
VIDEO_EXTS = {".mp4", ".webm", ".mov", ".avi", ".mkv", ".m4v"}
MEDIA_EXTS = IMAGE_EXTS | VIDEO_EXTS

# markdown 图片引用：![alt](path)
_MD_IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


@dataclass
class MediaFile:
    """zip 内提取出的媒体文件"""
    zip_path: str       # zip 内绝对路径（已修复编码、去掉开头 /）
    filename: str       # 原始文件名（basename）
    ext: str            # 扩展名（小写，含 .）
    content_type: str   # MIME 类型
    data: bytes         # 文件字节


def ext_of(name: str) -> str:
    i = name.rfind(".")
    return name[i:].lower() if i >= 0 else ""


def _decode_zip_filename(info: zipfile.ZipInfo) -> str:
    """修复 zip 内中文文件名乱码。

    macOS/Windows 部分 zip 工具用 UTF-8 存文件名但不设置 0x800 flag，
    Python zipfile 误按 cp437 解码导致乱码。未设置 flag 时尝试反向恢复。
    """
    name = info.filename
    if not (info.flag_bits & 0x800):
        try:
            name = name.encode("cp437").decode("utf-8")
        except (UnicodeDecodeError, UnicodeEncodeError):
            pass
    return name


def parse_pdf(data: bytes) -> str:
    """用 PyMuPDF 提取 PDF 文本；无文字层（扫描件）时回退 OCR。"""
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    parts: list[str] = []
    for page in doc:
        parts.append(page.get_text("text"))
    doc.close()
    text = "\n\n".join(parts)

    # 扫描件检测：文字层过少（<50 字符）且 OCR 可用时，逐页渲染 + OCR
    if len(text.strip()) < 50 and _ocr_available():
        ocr_text = _ocr_pdf(data)
        if ocr_text.strip():
            return ocr_text
    return text


def _ocr_available() -> bool:
    """检测系统是否安装 tesseract。"""
    return shutil.which("tesseract") is not None


def _ocr_pdf(data: bytes) -> str:
    """扫描件 PDF：逐页渲染成 PNG 并用 tesseract OCR 识别。"""
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    parts: list[str] = []
    try:
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            png = pix.tobytes("png")
            text = _ocr_png(png)
            if text.strip():
                parts.append(text.strip())
    finally:
        doc.close()
    return "\n\n".join(parts)


def _ocr_png(png: bytes) -> str:
    """调用 tesseract 识别单张 PNG（中文简体 + 英文）。"""
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp_path = tmp.name
    try:
        tmp.write(png)
        tmp.close()
        result = subprocess.run(
            ["tesseract", tmp_path, "stdout", "-l", "chi_sim+eng", "--psm", "6"],
            capture_output=True,
            text=True,
            timeout=60,
        )
        return result.stdout or ""
    except (OSError, subprocess.SubprocessError):
        return ""
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


def parse_image(data: bytes) -> str:
    """图片 OCR：用 PIL 统一转 PNG 后调用 tesseract 识别文字。"""
    if not _ocr_available():
        return ""
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(data))
        # 统一转 RGB（处理 RGBA/P/CMYK/LA 等模式，PNG 才兼容）
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return _ocr_png(buf.getvalue()).strip()
    except Exception:  # noqa: BLE001
        return ""


def parse_docx(data: bytes) -> str:
    """用 python-docx 提取 Word 文本"""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # 表格
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_xlsx(data: bytes) -> str:
    """用 openpyxl 提取 Excel 文本"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## 工作表：{ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def parse_zip(data: bytes) -> list[dict]:
    """解压 zip，返回其中可解析的文档列表 [{name, ext, text}]"""
    results: list[dict] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = _decode_zip_filename(info)
            ext = ext_of(name)
            raw = zf.read(info)
            try:
                if ext in TEXT_EXTS:
                    text = raw.decode("utf-8", errors="replace")
                    results.append({"name": name, "ext": ext, "text": text})
                elif ext == ".pdf":
                    results.append({"name": name, "ext": ext, "text": parse_pdf(raw)})
                elif ext == ".docx":
                    results.append({"name": name, "ext": ext, "text": parse_docx(raw)})
                elif ext == ".xlsx":
                    results.append({"name": name, "ext": ext, "text": parse_xlsx(raw)})
            except Exception:  # noqa: BLE001
                # 单个文件解析失败（如损坏的 PDF）不影响整个 zip
                continue
    return results


def extract_media(filename: str, data: bytes) -> list[MediaFile]:
    """提取 zip 内的图片/动图/视频文件；非 zip 返回空列表。"""
    if ext_of(filename) != ".zip":
        return []

    from app.services import media

    results: list[MediaFile] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = _decode_zip_filename(info)
            ext = ext_of(name)
            if ext not in MEDIA_EXTS:
                continue
            raw = zf.read(info)
            zip_path = name.lstrip("/")
            results.append(
                MediaFile(
                    zip_path=zip_path,
                    filename=name.rsplit("/", 1)[-1],
                    ext=ext,
                    content_type=media.content_type(ext),
                    data=raw,
                )
            )
    return results


def replace_md_image_refs(text: str, md_path: str, path_map: dict[str, str]) -> str:
    """把 markdown 里的图片引用路径替换为 /api/v1/media/{存储文件名}。

    md_path 为 md 文件在 zip 内的路径（已修复编码），path_map 为
    zip 内绝对路径 -> 存储文件名 的映射。跳过 http/https 外部链接，
    未命中映射的引用保持原样。
    """
    md_dir = posixpath.dirname(md_path)

    def repl(m: re.Match) -> str:
        alt = m.group(1)
        ref = m.group(2).strip()
        if not ref or ref.startswith(("http://", "https://")):
            return m.group(0)
        abs_path = posixpath.normpath(posixpath.join(md_dir, ref)).lstrip("/")
        stored = path_map.get(abs_path)
        if stored:
            return f"![{alt}](/api/v1/media/{stored})"
        return m.group(0)

    return _MD_IMG_RE.sub(repl, text)


def parse_file(filename: str, data: bytes) -> list[dict]:
    """解析单个文件，返回文档列表 [{name, ext, text}]"""
    ext = ext_of(filename)
    if ext == ".zip":
        return parse_zip(data)
    if ext in TEXT_EXTS:
        text = data.decode("utf-8", errors="replace")
        return [{"name": filename, "ext": ext, "text": text}]
    if ext == ".pdf":
        return [{"name": filename, "ext": ext, "text": parse_pdf(data)}]
    if ext == ".docx":
        return [{"name": filename, "ext": ext, "text": parse_docx(data)}]
    if ext == ".xlsx":
        return [{"name": filename, "ext": ext, "text": parse_xlsx(data)}]
    return []


def html_to_markdown(html: str, base_url: str = "") -> tuple[str, str]:
    """网页 HTML → (标题, Markdown 正文)。

    用 BeautifulSoup 清洗（去 script/style/nav/footer 等噪音，优先取
    article/main 正文），再用 html2text 转 Markdown。返回 (title, markdown)。
    """
    import bs4
    import html2text as _h2t

    soup = bs4.BeautifulSoup(html, "html.parser")
    title = (soup.title.get_text(strip=True) if soup.title else "") or ""

    # 移除噪音标签
    for tag in soup(["script", "style", "noscript", "iframe", "nav", "header", "footer", "aside", "form", "button"]):
        tag.decompose()

    # 正文：优先 article / main，否则 body
    body = soup.find("article") or soup.find("main") or soup.body or soup

    h = _h2t.HTML2Text()
    h.ignore_links = False
    h.ignore_images = True  # 网页图片链接通常不稳定，忽略
    h.body_width = 0  # 不折行，保持原样
    h.unicode_snob = True
    md = h.handle(str(body)).strip()
    return title, md
