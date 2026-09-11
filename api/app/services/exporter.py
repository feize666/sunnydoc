"""多格式导出服务：md / docx / pdf / html / json / zip。

PDF 中文方案：reportlab 内置 CID 字体 STSong-Light（UnicodeCIDFont），
无需下载/依赖系统字体即可正确渲染中文，避免乱码或方块。
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from typing import Any
from xml.sax.saxutils import escape as _xml_escape

from app.core.config import MEDIA_DIR

# markdown 标题层级 -> 正文/标题
_MD_H_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_LINK_RE = re.compile(r"!\[([^\]]*)\]\([^)]+\)|\[([^\]]+)\]\([^)]+\)")
_MD_LIST_RE = re.compile(r"^\s*(?:[-*+]|\d+\.)\s+")


def build(fmt: str, docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    """按格式导出文档列表，返回 (内容流, 文件名, Content-Type)。"""
    fmt = (fmt or "").strip().lower()
    handlers = {
        "md": _export_md,
        "markdown": _export_md,
        "docx": _export_docx,
        "pdf": _export_pdf,
        "html": _export_html,
        "json": _export_json,
        "zip": _export_zip,
    }
    handler = handlers.get(fmt)
    if handler is None:
        raise ValueError(f"不支持的导出格式：{fmt}")
    return handler(docs)


# ---------- 公共工具 ----------

def _safe_filename(name: str) -> str:
    """去除文件名中的非法字符。"""
    name = re.sub(r'[\\/:*?"<>|]', "_", name).strip()
    return name or "export"


def _export_filename(docs: list[dict[str, Any]], ext: str) -> str:
    if len(docs) == 1:
        return f"{_safe_filename(docs[0]['title'])}.{ext}"
    return f"export.{ext}"


def _md_plain(text: str) -> str:
    """把 markdown 语法尽量去掉，保留可读文本（用于 PDF 等纯文本渲染）。"""
    text = _MD_LINK_RE.sub(lambda m: m.group(1) or m.group(2) or "", text)
    text = _MD_LIST_RE.sub("", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = text.replace("~~", "").replace(">", "")
    return text.strip()


def _md_level(text: str) -> tuple[int, str]:
    """识别标题层级，返回 (level, 去掉 # 后的文本)。"""
    m = _MD_H_RE.match(text)
    if m:
        return min(len(m.group(1)), 3), m.group(2).strip()
    return 0, text


# ---------- markdown ----------

def _export_md(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    if len(docs) == 1:
        content = docs[0]["text"].encode("utf-8")
        filename = f"{_safe_filename(docs[0]['title'])}.md"
    else:
        parts = [f"# {d['title']}\n\n{d['text']}" for d in docs]
        content = "\n\n---\n\n".join(parts).encode("utf-8")
        filename = "export.md"
    return content, filename, "text/markdown; charset=utf-8"


# ---------- docx ----------

def _export_docx(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()

    def add_para(text: str, size: int, bold: bool) -> None:
        p = document.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for d in docs:
        add_para(d["title"], 18, True)
        for line in d["text"].split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            level, txt = _md_level(stripped)
            txt = _md_plain(txt)
            if not txt:
                continue
            if level == 1:
                add_para(txt, 16, True)
            elif level == 2:
                add_para(txt, 14, True)
            elif level == 3:
                add_para(txt, 12, True)
            else:
                add_para(txt, 11, False)

    buf = io.BytesIO()
    document.save(buf)
    return (
        buf.getvalue(),
        _export_filename(docs, "docx"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------- pdf ----------

def _export_pdf(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    # 内置中文 CID 字体，无需外部字体文件
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    styles = {
        0: ParagraphStyle(
            "body", fontName="STSong-Light", fontSize=11, leading=18,
            alignment=TA_LEFT,
        ),
        1: ParagraphStyle(
            "h1", fontName="STSong-Light", fontSize=20, leading=28, spaceAfter=12,
        ),
        2: ParagraphStyle(
            "h2", fontName="STSong-Light", fontSize=16, leading=24, spaceAfter=8,
        ),
        3: ParagraphStyle(
            "h3", fontName="STSong-Light", fontSize=13, leading=20, spaceAfter=6,
        ),
    }

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    story: list[Any] = []
    for d in docs:
        story.append(Paragraph(_xml_escape(_md_plain(d["title"])), styles[1]))
        for line in d["text"].split("\n"):
            stripped = line.rstrip()
            if not stripped.strip():
                continue
            level, txt = _md_level(stripped)
            txt = _md_plain(txt)
            if not txt:
                continue
            story.append(Paragraph(_xml_escape(txt), styles[level]))
            story.append(Spacer(1, 4))
    doc.build(story)

    return buf.getvalue(), _export_filename(docs, "pdf"), "application/pdf"


# ---------- html ----------

_CSS = """
body { font-family: -apple-system, "PingFang SC", "Microsoft YaHei", sans-serif;
       max-width: 860px; margin: 40px auto; padding: 0 24px; line-height: 1.7; color: #24292f; }
h1, h2, h3 { border-bottom: 1px solid #eaecef; padding-bottom: .3em; }
pre { background: #f6f8fa; padding: 16px; border-radius: 6px; overflow: auto; }
code { background: #f6f8fa; padding: .2em .4em; border-radius: 3px; }
blockquote { border-left: 4px solid #dfe2e5; margin: 0; padding: 0 1em; color: #6a737d; }
img { max-width: 100%; }
table { border-collapse: collapse; }
th, td { border: 1px solid #dfe2e5; padding: 6px 13px; }
"""


def _export_html(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    import markdown as md

    parts = [
        f"<h1>{_xml_escape(d['title'])}</h1>\n{md.markdown(d['text'], extensions=['tables', 'fenced_code'])}"
        for d in docs
    ]
    html = (
        "<!DOCTYPE html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<title>导出</title><style>{_CSS}</style></head><body>"
        + "\n".join(parts)
        + "</body></html>"
    )
    return html.encode("utf-8"), _export_filename(docs, "html"), "text/html; charset=utf-8"


# ---------- json ----------

def _export_json(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    data = [
        {
            "id": d["id"],
            "title": d["title"],
            "text": d["text"],
            "source": d["source"],
            "ext": d["ext"],
            "folder_id": d.get("folder_id"),
            "created_at": d["created_at"],
        }
        for d in docs
    ]
    content = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    return content, "export.json", "application/json; charset=utf-8"


# ---------- zip ----------

def _export_zip(docs: list[dict[str, Any]]) -> tuple[Any, str, str]:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        seen: set[str] = set()
        for d in docs:
            base = _safe_filename(d["title"])
            name = f"{base}.md"
            i = 1
            while name in seen:
                i += 1
                name = f"{base}({i}).md"
            seen.add(name)
            zf.writestr(name, d["text"])
        # 打包已存储的媒体文件
        if MEDIA_DIR.exists():
            for f in sorted(MEDIA_DIR.iterdir()):
                if f.is_file():
                    zf.write(f, f"media/{f.name}")
    return buf.getvalue(), "export.zip", "application/zip"


# ---------- 知识库整体导出（zip，保留目录结构） ----------

def build_kb_zip(
    kb_name: str,
    docs: list[dict[str, Any]],
    folders: list[dict[str, Any]],
) -> tuple[Any, str, str]:
    """把整个知识库导出为 zip，按文件夹层级组织，含 media 与 manifest。"""
    import posixpath as _pp

    base = _safe_filename(kb_name or "knowledge_base")

    # 构建 folder id -> 相对路径 的映射（递归，防御循环引用）
    folder_paths: dict[str, str] = {}
    _resolving: set[str] = set()

    def resolve(fid: str) -> str:
        if fid in folder_paths:
            return folder_paths[fid]
        if fid in _resolving:  # 防御循环引用
            return ""
        _resolving.add(fid)
        f = next((x for x in folders if x.get("id") == fid), None)
        path = ""
        if f:
            parent = f.get("parent_id")
            parent_path = resolve(parent) if parent else ""
            name = _safe_filename(f.get("name") or "未命名文件夹")
            path = _pp.join(parent_path, name) if parent_path else name
        _resolving.discard(fid)
        folder_paths[fid] = path
        return path

    for f in folders:
        resolve(f.get("id") or "")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        seen: set[str] = set()
        manifest: list[dict[str, Any]] = []
        for d in docs:
            subdir = folder_paths.get(d.get("folder_id") or "", "")
            stem = _safe_filename(d["title"])
            name = f"{stem}.md"
            full = _pp.join(base, subdir, name) if subdir else _pp.join(base, name)
            i = 1
            while full in seen:
                i += 1
                alt = f"{stem}({i}).md"
                full = _pp.join(base, subdir, alt) if subdir else _pp.join(base, alt)
            seen.add(full)
            zf.writestr(full, d["text"])
            manifest.append(
                {
                    "title": d["title"],
                    "path": full,
                    "source": d.get("source"),
                    "ext": d.get("ext"),
                }
            )

        # 媒体文件
        if MEDIA_DIR.exists():
            for f in sorted(MEDIA_DIR.iterdir()):
                if f.is_file():
                    zf.write(f, _pp.join(base, "media", f.name))

        # 清单
        zf.writestr(
            _pp.join(base, "manifest.json"),
            json.dumps(manifest, ensure_ascii=False, indent=2),
        )

    return buf.getvalue(), f"{base}.zip", "application/zip"
